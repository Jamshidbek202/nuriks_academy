from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path("/Users/muhammad/projects/nuriks_academy")
OUTPUT = ROOT / "docs" / "Nuriks_Academy_Finance_Decision_Questionnaire_RU.docx"
SKILL_ROOT = Path(
    "/Users/muhammad/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.715.12143/skills/documents"
)
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402


# standard_business_brief token map (plus named form/title overrides)
FONT = "Calibri"
INK = "0B2545"
ACCENT = "2E74B5"
ACCENT_DARK = "1F4D78"
MUTED = "5B6573"
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
ANSWER_FILL = "F8FAFC"
ANSWER_BORDER = "B8C6D8"
WHITE = "FFFFFF"
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_run_font(run, *, size=None, color=None, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "ru-RU")
    run._element.get_or_add_rPr().append(lang)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_borders(cell, *, color=ANSWER_BORDER, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_row_min_height(row, twips):
    tr_pr = row._tr.get_or_add_trPr()
    tr_height = OxmlElement("w:trHeight")
    tr_height.set(qn("w:val"), str(twips))
    tr_height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(tr_height)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="D6DCE5", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_no_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_paragraph_left_border(paragraph, color=ACCENT, size="18", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), space)
    left.set(qn("w:color"), color)
    p_bdr.append(left)


def add_page_field(paragraph):
    paragraph.add_run("Стр. ")
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    for item in paragraph.runs:
        set_run_font(item, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("20242A")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title.font.size = Pt(26)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(INK)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.line_spacing = 1.0

    subtitle = styles["Subtitle"]
    subtitle.font.name = FONT
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    subtitle.font.size = Pt(13)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.paragraph_format.line_spacing = 1.1

    for style_name, size, color, before, after in (
        ("Heading 1", 16, ACCENT, 16, 8),
        ("Heading 2", 13, ACCENT, 12, 6),
        ("Heading 3", 12, ACCENT_DARK, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def configure_sections(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header_p = section.header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_p.paragraph_format.space_after = Pt(0)
        header_p.text = "NURIKS ACADEMY  |  СОГЛАСОВАНИЕ ФИНАНСОВОЙ СИСТЕМЫ"
        for run in header_p.runs:
            set_run_font(run, size=8.5, color=MUTED, bold=True)

        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_p.paragraph_format.space_before = Pt(0)
        footer_p.paragraph_format.space_after = Pt(0)
        add_page_field(footer_p)


def add_labeled_paragraph(doc, label, text, *, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = keep
    p.paragraph_format.keep_with_next = keep
    label_run = p.add_run(label + " ")
    set_run_font(label_run, size=11, color=INK, bold=True)
    text_run = p.add_run(text)
    set_run_font(text_run, size=11, color="20242A")
    return p


def add_callout(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.10
    set_paragraph_shading(p, CALLOUT_FILL)
    set_paragraph_left_border(p)
    label_run = p.add_run(label + " ")
    set_run_font(label_run, size=10.5, color=ACCENT_DARK, bold=True)
    text_run = p.add_run(text)
    set_run_font(text_run, size=10.5, color="20242A")
    return p


def add_answer_box(doc, prompt):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run("Согласованное решение")
    set_run_font(run, size=10.5, color=INK, bold=True)

    table = doc.add_table(rows=1, cols=1)
    table.style = None
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_shading(cell, ANSWER_FILL)
    set_cell_borders(cell)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.10
    run = paragraph.add_run(prompt)
    set_run_font(run, size=9.5, color="788493", italic=True)
    blank = cell.add_paragraph("\n")
    blank.paragraph_format.space_after = Pt(0)
    set_row_min_height(table.rows[0], 820)
    apply_table_geometry(
        table,
        [CONTENT_DXA],
        table_width_dxa=CONTENT_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa={"top": 100, "bottom": 100, "start": 120, "end": 120},
    )

    sign = doc.add_paragraph()
    sign.paragraph_format.space_before = Pt(3)
    sign.paragraph_format.space_after = Pt(8)
    run = sign.add_run(
        "Ответственный: ____________________    Дата: __________    "
        "Статус: Согласовано / Открыто"
    )
    set_run_font(run, size=9, color=MUTED)


def add_data_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = None
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for idx, value in enumerate(headers):
        cell = header_row.cells[idx]
        set_cell_shading(cell, LIGHT_FILL)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(value)
        set_run_font(run, size=9.5, color=INK, bold=True)
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run_font(run, size=9.5, color="20242A")
    set_table_borders(table)
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=CONTENT_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa=CELL_MARGINS,
    )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


QUESTIONS = [
    {
        "section": "1. Расчёт стоимости обучения и уроков",
        "number": 1,
        "title": "Что именно означают утверждённые цены",
        "question": (
            "Указанные суммы являются фиксированной месячной платой или ценой за "
            "определённое стандартное количество уроков? Нужно решить, как получить "
            "цену одного урока и что происходит в месяце с праздниками."
        ),
        "why": (
            "Если каждый месяц делить тариф на фактическое число уроков этого месяца, "
            "исключение праздника может не уменьшить счёт. Если делить на установленную "
            "норму, итоговая сумма действительно зависит от числа оплачиваемых уроков."
        ),
        "recommendation": (
            "Зафиксировать для каждого тарифа стандартное число уроков. Цена урока = "
            "месячный тариф / стандартное число уроков. Счёт = сумма цен всех "
            "оплачиваемых уроков по действовавшим на их даты тарифам."
        ),
        "prompt": (
            "Укажите: фиксированный месяц или норма уроков; формулу цены урока; что "
            "происходит при меньшем/большем количестве уроков."
        ),
    },
    {
        "section": "1. Расчёт стоимости обучения и уроков",
        "number": 2,
        "title": "Стандартное количество уроков и обязательное расписание",
        "question": (
            "Сколько уроков включает каждый тариф: например, 12 уроков в месяц или "
            "другое число? Могут ли группы учиться по разной частоте? Какие поля "
            "обязательны в расписании: дни недели, время, длительность, преподаватель, "
            "дата начала и дата окончания?"
        ),
        "why": (
            "Без точного расписания невозможно заранее построить календарь уроков, "
            "исключить праздники, рассчитать счёт и разделить оплату преподавателей."
        ),
        "recommendation": (
            "Сделать расписание источником календаря уроков. Хранить дни недели, "
            "время начала, длительность, преподавателя и период действия расписания. "
            "Изменения расписания также должны иметь дату вступления в силу."
        ),
        "prompt": (
            "Укажите норму уроков для каждого формата, допустимую частоту занятий, "
            "длительность урока и обязательные поля расписания."
        ),
    },
    {
        "section": "1. Расчёт стоимости обучения и уроков",
        "number": 3,
        "title": "Какие уроки оплачиваются, а какие исключаются",
        "question": (
            "Нужно утвердить правила для проведённого урока, отсутствия ученика, "
            "официального выходного, закрытия центра, отмены преподавателем, переноса "
            "и дополнительного/замещающего урока."
        ),
        "why": (
            "Каждый статус влияет одновременно на счёт ученика, долг, доход центра и "
            "заработок преподавателя. Неясный статус создаст двойное начисление или "
            "неправильное исключение."
        ),
        "recommendation": (
            "Проведено — начислить и оплатить преподавателю. Ученик отсутствовал — "
            "начислить и оплатить преподавателю. Праздник/закрытие центра — не начислять. "
            "Отмена преподавателем — не начислять, пока не проведён заменяющий урок. "
            "Заменяющий урок не должен начисляться дважды."
        ),
        "prompt": (
            "Для каждого статуса укажите: начислять ученику да/нет; начислять "
            "преподавателю да/нет; нужен ли перенос; кто вправе изменить статус."
        ),
    },
    {
        "section": "2. Доля и заработок преподавателей",
        "number": 4,
        "title": "Когда возникает заработок преподавателя",
        "question": (
            "Преподаватель получает свою долю после выставления счёта ученику или "
            "только после фактического получения денег центром? Например, ученик должен "
            "450 000 сум, но ещё не заплатил."
        ),
        "why": (
            "Это определяет задолженность центра перед преподавателем и риск оплаты "
            "зарплаты из денег, которые ещё не поступили."
        ),
        "recommendation": (
            "Отдельно показывать заработано по проведённым оплачиваемым урокам, "
            "покрыто фактическими платежами и выплачено преподавателю. Руководство "
            "должно отдельно решить, разрешена ли выплата до получения денег."
        ),
        "prompt": (
            "Выберите момент возникновения заработка и момент, когда сумму разрешено "
            "выплатить. Укажите правило для частично оплаченных счетов."
        ),
    },
    {
        "section": "2. Доля и заработок преподавателей",
        "number": 5,
        "title": "База расчёта 40% и 50%",
        "question": (
            "Подтвердите: преподаватель получает 40% от всей выручки всех учеников "
            "обычной/мини-группы и 50% по индивидуальному обучению. Доля считается до "
            "или после скидок, возвратов, кредитов и списаний? Что происходит при смене "
            "преподавателя в середине месяца?"
        ),
        "why": (
            "Слово «выручка» может означать выставленную, полученную или чистую сумму. "
            "Для зарплаты нужна одна точная формула и историческая привязка к урокам."
        ),
        "recommendation": (
            "Для каждого урока связывать сумму ученика, применённую скидку, чистую "
            "начисленную сумму и преподавателя этого урока. Долю считать от чистой "
            "суммы после утверждённых скидок/корректировок."
        ),
        "prompt": (
            "Подтвердите проценты; определите базу процента; правило скидок/возвратов; "
            "распределение при смене преподавателя."
        ),
    },
    {
        "section": "3. Тарифы, программы и изменение групп",
        "number": 6,
        "title": "Дата вступления в силу изменений",
        "question": (
            "Как применять изменение цены, формата группы, расписания или процента "
            "преподавателя в середине месяца? Разрешены ли изменения задним числом?"
        ),
        "why": (
            "Если система перезапишет старое значение, ранее проведённые уроки, счета "
            "и зарплата изменятся без объяснения."
        ),
        "recommendation": (
            "Каждое изменение имеет дату вступления в силу. Старые уроки сохраняют "
            "старые условия, новые уроки используют новые. После финализации заднее "
            "число запрещено; исправления оформляются отдельной корректировкой."
        ),
        "prompt": (
            "Укажите минимально допустимую дату изменения; кто может менять; нужны ли "
            "уведомление и подтверждение; допускается ли заднее число до финализации."
        ),
    },
    {
        "section": "3. Тарифы, программы и изменение групп",
        "number": 7,
        "title": "Тариф Pre-IELTS",
        "question": (
            "Должен ли Pre-IELTS иметь те же цены, что IELTS, для обычной группы, "
            "мини-группы и индивидуального обучения? Должны ли IELTS и Pre-IELTS "
            "оставаться отдельными названиями программ?"
        ),
        "why": (
            "В исходных ценах IELTS и Pre-IELTS объединены только для обычной группы. "
            "Для остальных форматов правило нужно зафиксировать."
        ),
        "recommendation": (
            "Хранить IELTS и Pre-IELTS как отдельные программы, но назначить им "
            "одинаковые стартовые цены: 550 000 / 750 000 / 1 400 000 сум."
        ),
        "prompt": (
            "Подтвердите цены Pre-IELTS для трёх форматов и необходимость отдельного "
            "названия программы в отчётах."
        ),
    },
    {
        "section": "3. Тарифы, программы и изменение групп",
        "number": 8,
        "title": "Размер обычной, мини- и индивидуальной группы",
        "question": (
            "Сколько учеников разрешено в обычной группе и мини-группе? Должна ли "
            "система только предупреждать или запрещать превышение? Индивидуальная "
            "группа всегда содержит ровно одного ученика?"
        ),
        "why": (
            "Размер группы влияет на формат, цену и доход. Без лимитов группа может "
            "сохранить более дешёвый/дорогой формат неправильно."
        ),
        "recommendation": (
            "Установить явные минимумы и максимумы. При нарушении показывать блокирующее "
            "предупреждение; изменение формата выполнять только с датой вступления в силу."
        ),
        "prompt": (
            "Укажите минимум/максимум для каждого формата и поведение системы при "
            "выходе за лимит."
        ),
    },
    {
        "section": "3. Тарифы, программы и изменение групп",
        "number": 9,
        "title": "Зачисление, уход, паузы, переводы и индивидуальные скидки",
        "question": (
            "Как начислять ученику, который пришёл или ушёл в середине месяца? Нужны ли "
            "официальная пауза, заморозка, перевод между группами, стипендия, скидка или "
            "индивидуальная цена?"
        ),
        "why": (
            "Эти события меняют число оплачиваемых уроков и не должны вручную стирать "
            "историю или создавать двойное начисление."
        ),
        "recommendation": (
            "У каждого события есть дата действия. Начислять только уроки в активном "
            "периоде зачисления. Скидки и индивидуальные цены хранить как отдельные "
            "версии с причиной, сроком и автором."
        ),
        "prompt": (
            "Опишите правила начала/окончания, паузы, перевода, скидок и необходимые "
            "подтверждающие причины."
        ),
    },
    {
        "section": "4. Счета, задолженность и платежи",
        "number": 10,
        "title": "Формирование, проверка и срок оплаты счёта",
        "question": (
            "Система должна автоматически создавать счёт за предыдущий месяц 1-го "
            "числа или сначала формировать черновик для проверки? Какой срок оплаты и "
            "с какого дня долг считается просроченным?"
        ),
        "why": (
            "Автоматически окончательно проведённая ошибка сразу повлияет на долги, "
            "напоминания, отчёты и зарплаты."
        ),
        "recommendation": (
            "1-го числа по времени Ташкента автоматически создавать черновики. "
            "Менеджер/суперадмин проверяет предупреждения и финализирует. После "
            "финализации изменения идут только корректировками."
        ),
        "prompt": (
            "Укажите: черновик или автоматическая финализация; ответственный; срок "
            "оплаты; день просрочки; действие, если счёт не проверен вовремя."
        ),
    },
    {
        "section": "4. Счета, задолженность и платежи",
        "number": 11,
        "title": "Исправления и неизменяемый журнал",
        "question": (
            "Можно ли редактировать или удалять финализированный счёт, платёж, расход "
            "или начисление зарплаты? Каким образом исправляется ошибка?"
        ),
        "why": (
            "Удаление финансовой записи уничтожает связь между кассой, долгом, "
            "отчётностью и ответственным пользователем."
        ),
        "recommendation": (
            "Запретить удаление и тихое изменение финализированных записей. Использовать "
            "аннулирование, обратную проводку, кредит или корректировку с причиной, "
            "автором, временем и ссылкой на исходную запись."
        ),
        "prompt": (
            "Утвердите допустимые способы исправления и роли, которые могут проводить "
            "каждый способ."
        ),
    },
    {
        "section": "4. Счета, задолженность и платежи",
        "number": 12,
        "title": "Правила приёма наличных платежей",
        "question": (
            "Разрешены ли частичные платежи и переплата? Должна ли переплата стать "
            "авансом? К какому счёту привязывать платёж: к выбранному или к самому "
            "старому долгу? Нужен ли номер квитанции? Кто может отменить ошибочный платёж?"
        ),
        "why": (
            "Ресепшен будет принимать наличные. Без правил один платёж может неверно "
            "закрыть долг или исчезнуть из кассы."
        ),
        "recommendation": (
            "Разрешить частичный платёж и аванс. По умолчанию закрывать самый старый "
            "долг, но показывать распределение до подтверждения. Каждому платежу давать "
            "уникальную квитанцию; отмена — только обратной записью руководителя."
        ),
        "prompt": (
            "Укажите правила частичной оплаты, аванса, распределения, квитанции, "
            "отмены и обязательных комментариев. Карта пока не включается."
        ),
    },
    {
        "section": "4. Счета, задолженность и платежи",
        "number": 13,
        "title": "Округление сумм и распределение остатка",
        "question": (
            "Как округлять цену урока и долю преподавателя, если деление даёт дробную "
            "сумму? Кому относится остаток в 1 сум?"
        ),
        "why": (
            "Деньги должны храниться целыми сумами. Дробные числа и разное округление "
            "на фронтенде и бэкенде создадут расхождения."
        ),
        "recommendation": (
            "Хранить суммы целыми числами, никогда не использовать числа с плавающей "
            "точкой. Применять одно правило округления на сервере; остаток относить "
            "центру, чтобы доли всегда точно равнялись общей сумме."
        ),
        "prompt": (
            "Укажите метод округления и получателя остатка. Подтвердите, что в интерфейсе "
            "показываются только целые суммы."
        ),
    },
    {
        "section": "5. Расходы и зарплата персонала",
        "number": 14,
        "title": "Оплата бухгалтера, ресепшена, помощников и другого персонала",
        "question": (
            "Как оплачиваются бухгалтер, будущий ресепшен, support teachers и другие "
            "сотрудники: фиксированная месячная сумма, за урок, процент или индивидуальная "
            "схема? Бухгалтерские 500 000 сум — зарплата или услуга/расход?"
        ),
        "why": (
            "Разные схемы оплаты требуют разных оснований и дат начисления. Их нельзя "
            "смешивать с долей основных преподавателей без явной классификации."
        ),
        "recommendation": (
            "Разрешить настраиваемую схему для каждого сотрудника с периодом действия: "
            "фиксировано, за урок, процент или ручное начисление. Отдельно классифицировать "
            "сотрудника и внешнюю услугу."
        ),
        "prompt": (
            "Для каждой текущей/будущей должности укажите схему, сумму/процент, дату "
            "начисления и кто подтверждает выплату."
        ),
    },
    {
        "section": "5. Расходы и зарплата персонала",
        "number": 15,
        "title": "Плановый расход, задолженность и фактическая оплата",
        "question": (
            "Аренда 10 500 000, бухгалтер 500 000 и Wi-Fi 400 000 сум должны "
            "автоматически считаться оплаченными или сначала появляться как ожидаемые/к "
            "оплате? Как заносить переменные суммы электричества и газа?"
        ),
        "why": (
            "Созданный расход и реально выданные деньги — разные события. Иначе отчёт "
            "о денежных средствах будет неточным."
        ),
        "recommendation": (
            "Шаблон регулярного расхода создаёт обязательство «к оплате». Менеджер или "
            "суперадмин отдельно фиксирует фактическую дату, сумму, способ и примечание. "
            "Электричество/газ создаются ежемесячно с ручной фактической суммой."
        ),
        "prompt": (
            "Для каждого расхода укажите дату начисления, срок оплаты, постоянная/" 
            "переменная сумма, подтверждающий документ и кто отмечает оплату."
        ),
    },
    {
        "section": "6. Роли и доступ к финансовым данным",
        "number": 16,
        "title": "Разделение полномочий суперадмина и менеджера",
        "question": (
            "Кто может менять глобальные цены, проценты преподавателей, регулярные "
            "расходы, финализировать счета, проводить зарплату и делать обратные записи?"
        ),
        "why": (
            "Одинаковый полный доступ у нескольких ролей повышает риск случайной или "
            "необъяснимой финансовой корректировки."
        ),
        "recommendation": (
            "Менеджер ведёт ежедневные операции и финализирует обычные счета. Только "
            "суперадмин меняет глобальные тарифы/проценты и отменяет финализированные "
            "операции. Все действия записываются в журнал."
        ),
        "prompt": (
            "Для каждой операции назначьте: просмотр, создание, подтверждение, "
            "корректировка и отмена — отдельно для суперадмина и менеджера."
        ),
    },
    {
        "section": "6. Роли и доступ к финансовым данным",
        "number": 17,
        "title": "Что видит и делает ресепшен",
        "question": (
            "Должен ли ресепшен видеть общую выручку, расходы, прибыль и зарплаты "
            "преподавателей? Подтвердите доступ к лидам, конвертации учеников, контактам, "
            "счетам, долгам, напоминаниям и регистрации наличной оплаты."
        ),
        "why": (
            "Ресепшену нужна информация для звонков и приёма денег, но не вся "
            "конфиденциальная финансовая картина центра."
        ),
        "recommendation": (
            "Ресепшен видит только учеников/родителей, счета, сроки, остатки, историю "
            "платежей, напоминания и свою кассу. Не видит зарплаты, расходы, общую "
            "прибыль и не меняет цены, скидки, счета или платёжную политику."
        ),
        "prompt": (
            "Перечислите разрешённые экраны и действия. Отдельно укажите, может ли "
            "ресепшен видеть точную сумму долга и историю платежей."
        ),
    },
    {
        "section": "6. Роли и доступ к финансовым данным",
        "number": 18,
        "title": "Временная учётная запись reception",
        "question": (
            "Текущий вход поддерживает логин reception и пароль Reception@2025? Нужно "
            "ли требовать смену пароля при первом входе? Как эта роль будет безопасно "
            "воссоздана после будущего перехода на вход по номеру телефона и очистки базы?"
        ),
        "why": (
            "Временная авторизация не должна смешаться с будущей моделью и потерять "
            "ограничения роли после повторного заполнения базы."
        ),
        "recommendation": (
            "Создать повторяемую seed-запись роли и временного пользователя без "
            "дублирования. Хранить пароль только в защищённом хеше и предусмотреть "
            "безопасную миграцию/замену на телефонную авторизацию."
        ),
        "prompt": (
            "Подтвердите логин, временный пароль, смену при первом входе, будущий номер "
            "телефона и владельца решения о повторном создании после очистки базы."
        ),
    },
    {
        "section": "7. Напоминания и отношения ученик-родитель",
        "number": 19,
        "title": "Канал и график уведомлений",
        "question": (
            "Что реализовать сейчас: внутренний список звонков для ресепшена, SMS, "
            "Telegram или другой канал? За сколько дней до срока и через сколько дней "
            "после просрочки отправлять напоминания?"
        ),
        "why": (
            "Уведомление без утверждённого получателя, канала и расписания может уйти "
            "не тому человеку или с неверной суммой."
        ),
        "recommendation": (
            "На первом этапе создать очередь задач/звонков с точной суммой, сроком, "
            "контактом и историей попыток. Внешние SMS/Telegram подключать после "
            "утверждения текста, согласий и поставщика."
        ),
        "prompt": (
            "Укажите текущий канал, дни напоминаний, рабочее время, язык сообщения, "
            "ответственного и условия остановки напоминаний."
        ),
    },
    {
        "section": "7. Напоминания и отношения ученик-родитель",
        "number": 20,
        "title": "Кто отвечает за оплату и получает напоминание",
        "question": (
            "Может ли у ученика быть несколько родителей/контактов? Может ли один "
            "родитель оплачивать нескольких учеников? Кто является основным плательщиком "
            "и кто получает счёт/напоминание?"
        ),
        "why": (
            "Контакт ученика и финансово ответственный человек могут отличаться. "
            "Неправильная связь приведёт к раскрытию долга или пропущенному звонку."
        ),
        "recommendation": (
            "Разделить ученика, контактное лицо и плательщика. Разрешить несколько "
            "контактов с флагами: основной, получает счета, получает напоминания, "
            "разрешён доступ к финансовой информации."
        ),
        "prompt": (
            "Опишите допустимые связи, правила выбора основного плательщика и кому "
            "можно сообщать сумму долга."
        ),
    },
    {
        "section": "8. Отчёты, данные и будущая миграция",
        "number": 21,
        "title": "Обязательные финансовые отчёты",
        "question": (
            "Какие показатели нужны на одной финансовой странице и в выгрузках: "
            "выставленная выручка, полученные деньги, долги, просрочка, расходы, зарплата, "
            "денежный поток, прибыль, разрез программы/группы/преподавателя?"
        ),
        "why": (
            "Полученные деньги и заработанная выручка — не одно и то же. Один общий "
            "показатель может скрыть долги или неоплаченные расходы."
        ),
        "recommendation": (
            "Показывать отдельно начислено, получено, к получению, просрочено, расходы "
            "начислены/оплачены, зарплата заработана/выплачена, денежный поток и прибыль. "
            "Добавить фильтры и журнал операций."
        ),
        "prompt": (
            "Отметьте обязательные показатели, фильтры, периоды и необходимость Excel, "
            "PDF, печатной квитанции и экспорта для бухгалтера."
        ),
    },
    {
        "section": "8. Отчёты, данные и будущая миграция",
        "number": 22,
        "title": "Существующие данные, очистка базы и повторное заполнение",
        "question": (
            "Нужно ли перенести текущие группы/учеников в новые форматы и цены сейчас? "
            "Что обязательно должно безопасно восстановиться после будущей очистки базы "
            "из-за телефонной авторизации?"
        ),
        "why": (
            "Логика финансов не должна зависеть от ручного повторного ввода или "
            "исчезнуть вместе с тестовыми пользователями."
        ),
        "recommendation": (
            "Разделить миграцию данных и seed системных настроек. Повторно создавать "
            "роли, тарифы, проценты и шаблоны расходов безопасным, идемпотентным "
            "процессом; финансовую историю не уничтожать без отдельного экспорта/архива."
        ),
        "prompt": (
            "Укажите, что мигрировать сейчас, что будет очищено позже, какие данные "
            "экспортировать и кто разрешает окончательную очистку."
        ),
    },
    {
        "section": "9. Дополнительные финансовые гарантии",
        "number": 23,
        "title": "Рабочий часовой пояс и календарь центра",
        "question": (
            "Подтвердите, что все даты, месяцы, счета и задания запускаются по часовому "
            "поясу Asia/Tashkent. Кто ведёт официальный календарь праздников и закрытий?"
        ),
        "why": (
            "Сервер или пользователь могут находиться в другом часовом поясе. Ошибка на "
            "границе месяца применит тариф или создаст счёт не в тот день."
        ),
        "recommendation": (
            "Хранить точные моменты времени в UTC, а бизнес-даты рассчитывать только по "
            "Asia/Tashkent. Праздники/закрытия ведутся централизованно с причиной и автором."
        ),
        "prompt": (
            "Подтвердите часовой пояс, рабочую неделю, источник праздников, кто может "
            "добавлять закрытие и как уведомлять затронутые группы."
        ),
    },
    {
        "section": "9. Дополнительные финансовые гарантии",
        "number": 24,
        "title": "Касса и ежедневная сверка наличных",
        "question": (
            "Нужна ли отдельная кассовая смена ресепшена: начальный остаток, все "
            "принятые деньги, возвраты, передача менеджеру, фактический остаток и "
            "расхождение? Есть ли одна касса/филиал или несколько?"
        ),
        "why": (
            "Записанный платёж не доказывает, что наличные физически находятся в кассе "
            "или были переданы ответственному лицу."
        ),
        "recommendation": (
            "Ввести ежедневное открытие/закрытие кассы, неизменяемый журнал, уникальные "
            "квитанции и подтверждение расхождения менеджером. На первом этапе — одна "
            "касса, если центр действительно один."
        ),
        "prompt": (
            "Укажите число касс/филиалов, начальный остаток, процесс передачи, допустимое "
            "расхождение и кто закрывает/подтверждает смену."
        ),
    },
    {
        "section": "9. Дополнительные финансовые гарантии",
        "number": 25,
        "title": "Возвраты, кредиты, списание долга и льготы",
        "question": (
            "В каких случаях разрешены возврат денег, кредит на будущий период, списание "
            "долга, бесплатный урок или льгота? Кто утверждает и как это влияет на долю "
            "преподавателя?"
        ),
        "why": (
            "Эти операции уменьшают долг или деньги, но не должны маскироваться как "
            "удаление исходного платежа/счёта."
        ),
        "recommendation": (
            "Каждый случай оформлять отдельной типизированной записью с причиной, "
            "утверждающим лицом и ссылкой на исходный счёт/платёж. Автоматически "
            "пересчитывать связанное начисление преподавателю по утверждённому правилу."
        ),
        "prompt": (
            "Перечислите допустимые случаи, лимиты, утверждающие роли, подтверждающие "
            "документы и влияние на зарплату преподавателя."
        ),
    },
    {
        "section": "9. Дополнительные финансовые гарантии",
        "number": 26,
        "title": "Налоги и границы бухгалтерского учёта",
        "question": (
            "Должна ли система рассчитывать налоги/НДС и заменять официальный "
            "бухгалтерский учёт, или это управленческий учёт центра с экспортом данных "
            "бухгалтеру? Включены ли налоги в указанные цены?"
        ),
        "why": (
            "Налоговые правила имеют юридические последствия. Нельзя считать прибыль "
            "или сумму преподавателя правильно без понимания, включён ли налог."
        ),
        "recommendation": (
            "На первом этапе считать систему управленческим учётом и передавать "
            "бухгалтеру полный экспорт. Налоговые расчёты внедрять только после "
            "письменно утверждённых требований квалифицированного бухгалтера."
        ),
        "prompt": (
            "Укажите налоговый статус, включение налогов в цены, необходимые отчёты и "
            "имя бухгалтера, который подтвердит требования."
        ),
    },
    {
        "section": "9. Дополнительные финансовые гарантии",
        "number": 27,
        "title": "Кто окончательно утверждает формулы и запуск",
        "question": (
            "Кто имеет право окончательно утвердить тарифы, формулы уроков, зарплату, "
            "права доступа и результаты тестового расчёта перед использованием реальных "
            "денег?"
        ),
        "why": (
            "Даже технически правильная реализация может не соответствовать договорённости "
            "центра. Нужен один подписанный источник решений."
        ),
        "recommendation": (
            "Назначить владельца финансовой политики. До запуска провести контрольный "
            "месяц на тестовых данных и письменно сравнить ожидаемые счета, долги, кассу, "
            "расходы и зарплаты с результатом системы."
        ),
        "prompt": (
            "Укажите владельца политики, обязательных согласующих, критерии успешного "
            "тестового месяца и лицо, разрешающее запуск."
        ),
    },
    {
        "section": "10. Все входящие деньги",
        "number": 28,
        "title": "Доходы, не связанные с оплатой обучения",
        "question": (
            "Может ли центр получать другие деньги: регистрационный сбор, продажа "
            "материалов, аренда помещения, возврат от поставщика, пожертвование или "
            "другой доход? Нужна ли всегда категория «Другой доход»?"
        ),
        "why": (
            "Требование охватывает все входящие и исходящие деньги. Иначе касса может "
            "увеличиться без объяснимого источника."
        ),
        "recommendation": (
            "Добавить настраиваемые категории дохода и обязательную категорию «Другой "
            "доход». Хранить источник, плательщика, дату, сумму, способ, комментарий и "
            "подтверждающий документ."
        ),
        "prompt": (
            "Перечислите текущие виды другого дохода, кто их регистрирует и кто "
            "подтверждает."
        ),
    },
]


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(44)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("NURIKS ACADEMY")
    set_run_font(run, size=11, color=ACCENT, bold=True)

    title = doc.add_paragraph(style="Title")
    title.add_run("Анкета для согласования\nфинансовой системы")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run(
        "Рабочий документ для обсуждения с владельцем, менеджером, бухгалтером "
        "и представителями преподавателей"
    )

    callout = doc.add_paragraph()
    callout.paragraph_format.space_before = Pt(8)
    callout.paragraph_format.space_after = Pt(18)
    callout.paragraph_format.left_indent = Inches(0.18)
    set_paragraph_shading(callout, CALLOUT_FILL)
    set_paragraph_left_border(callout, color=ACCENT)
    run = callout.add_run(
        "Статус: обсуждение до разработки. Приложение, база данных и финансовые "
        "настройки не изменялись. Рекомендации внутри документа являются предложениями, "
        "а не утверждёнными правилами."
    )
    set_run_font(run, size=10.5, color=INK, bold=True)

    for label, value in (
        ("Дата подготовки:", "21 июля 2026 года"),
        ("Версия:", "1.0 — русская версия для согласования"),
        ("Участники обсуждения:", "____________________________________________"),
        ("Дата встречи:", "____________________________________________"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(label + " ")
        set_run_font(r1, size=10.5, color=INK, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5, color="20242A")

    doc.add_page_break()


def add_scope_and_known_values(doc):
    doc.add_heading("Как заполнять документ", level=1)
    add_labeled_paragraph(
        doc,
        "Шаг 1.",
        "Обсудите каждый вопрос с людьми, которые принимают финансовые решения.",
    )
    add_labeled_paragraph(
        doc,
        "Шаг 2.",
        "Запишите одно согласованное решение в серое поле. Если решения пока нет, "
        "поставьте статус «Открыто» и назначьте ответственного.",
    )
    add_labeled_paragraph(
        doc,
        "Шаг 3.",
        "Не оставляйте формулу подразумеваемой. Укажите даты, исключения, роли и "
        "примеры, особенно для середины месяца.",
    )
    add_labeled_paragraph(
        doc,
        "Шаг 4.",
        "После ответов заполните контрольные сценарии и финальное согласование в конце.",
    )

    doc.add_heading("Исходные цены, предоставленные для обсуждения", level=1)
    add_data_table(
        doc,
        ["Программа", "Обычная группа", "Мини-группа", "Индивидуально"],
        [
            ["General", "450 000 сум", "650 000 сум", "1 200 000 сум"],
            ["IELTS / Pre-IELTS", "550 000 сум", "750 000 сум", "1 400 000 сум"],
        ],
        [2520, 2280, 2280, 2280],
    )
    add_callout(
        doc,
        "Текущая доля преподавателя:",
        "40% по обычным и мини-группам; 50% по индивидуальным занятиям. "
        "Точная база и момент возникновения заработка должны быть утверждены в вопросах 4–5.",
    )

    doc.add_heading("Исходные расходы", level=1)
    add_data_table(
        doc,
        ["Категория", "Исходная сумма / правило"],
        [
            ["Аренда", "10 500 000 сум в месяц"],
            ["Бухгалтер", "500 000 сум в месяц"],
            ["Wi-Fi", "400 000 сум в месяц"],
            ["Электричество", "Сумма требует ежемесячного ввода"],
            ["Газ", "Сумма требует ежемесячного ввода"],
            ["Другой расход", "Обязательная категория для нерегулярных расходов"],
        ],
        [2700, 6660],
    )

    doc.add_heading("Охват будущей системы", level=1)
    add_labeled_paragraph(
        doc,
        "Финансы:",
        "одна страница для суперадмина и менеджера: начисления, поступления, долги, "
        "расходы, зарплаты, касса и отчёты.",
    )
    add_labeled_paragraph(
        doc,
        "Ресепшен:",
        "лиды, конвертация учеников, звонки по оплате, просмотр разрешённой информации "
        "и регистрация наличных платежей без управления тарифами.",
    )
    add_labeled_paragraph(
        doc,
        "Динамика:",
        "версионные цены, расписания, формат группы и кадровые изменения с датой "
        "вступления в силу; история не переписывается.",
    )
    add_labeled_paragraph(
        doc,
        "Будущее:",
        "переход на вход по номеру телефона, повторное заполнение базы, карточные "
        "платежи и внешние уведомления — без потери утверждённой финансовой логики.",
    )


def add_questions(doc):
    current_section = None
    for question in QUESTIONS:
        if question["section"] != current_section:
            if current_section is not None:
                doc.add_page_break()
            current_section = question["section"]
            doc.add_heading(current_section, level=1)
            p = doc.add_paragraph(
                "Заполните каждое решение в этой секции. Если ответа пока нет, "
                "назначьте ответственного и дату возврата к вопросу."
            )
            p.paragraph_format.space_after = Pt(10)
            for run in p.runs:
                set_run_font(run, size=10, color=MUTED, italic=True)

        # Named pagination override: keep the final approval question coherent.
        if question["number"] == 27:
            doc.add_page_break()

        heading = doc.add_heading(
            f"{question['number']}. {question['title']}", level=2
        )
        heading.paragraph_format.keep_with_next = True

        add_labeled_paragraph(
            doc, "Вопрос:", question["question"], keep=True
        )
        add_labeled_paragraph(
            doc, "Почему это важно:", question["why"], keep=False
        )
        add_callout(
            doc, "Рекомендуемый исходный вариант:", question["recommendation"]
        )
        add_answer_box(doc, question["prompt"])


def add_validation_and_signoff(doc):
    doc.add_page_break()
    doc.add_heading("Контрольные сценарии перед запуском", level=1)
    intro = doc.add_paragraph(
        "После реализации для каждого сценария необходимо заранее записать ожидаемый "
        "результат, затем сравнить его с результатом системы. Запуск с реальными деньгами "
        "не рекомендуется, пока обязательные сценарии не пройдены."
    )
    intro.paragraph_format.space_after = Pt(10)

    scenarios = [
        ("1", "Цена группы меняется 15-го числа месяца"),
        ("2", "Обычная группа становится мини-группой в середине месяца"),
        ("3", "Ученик отсутствует, но урок проводится"),
        ("4", "Официальный праздник исключает запланированный урок"),
        ("5", "Преподаватель отменяет урок, затем проводит замену"),
        ("6", "Ученик приходит, уходит, переводится или ставит обучение на паузу"),
        ("7", "Ученик не платит, платит частично или переплачивает"),
        ("8", "Наличный платёж записан ошибочно и проводится обратная запись"),
        ("9", "Преподаватель меняется в середине месяца"),
        ("10", "Электричество отличается от ожидаемой суммы"),
        ("11", "Касса при закрытии не совпадает с зарегистрированными платежами"),
        ("12", "База очищена и системные роли/тарифы создаются повторно"),
    ]
    table = add_data_table(
        doc,
        ["№", "Сценарий", "Ожидаемый результат записан", "Тест пройден"],
        [[num, text, "Да / Нет", "Да / Нет"] for num, text in scenarios],
        [540, 5580, 1800, 1440],
    )
    for row in table.rows[1:]:
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Нерешённые вопросы", level=1)
    add_answer_box(
        doc,
        "Перечислите номера вопросов со статусом «Открыто», ответственных и крайние "
        "даты принятия решения."
    )

    doc.add_page_break()
    doc.add_heading("Финальное согласование", level=1)
    add_callout(
        doc,
        "Важно:",
        "Подписание этой страницы означает согласование бизнес-логики для проектирования "
        "и тестирования. Это не заменяет юридическую или бухгалтерскую консультацию."
    )
    add_data_table(
        doc,
        ["Роль", "Ф.И.О.", "Решение", "Дата / подпись"],
        [
            ["Владелец финансовой политики", "", "Согласовано / Есть замечания", ""],
            ["Суперадмин", "", "Согласовано / Есть замечания", ""],
            ["Менеджер", "", "Согласовано / Есть замечания", ""],
            ["Бухгалтер", "", "Согласовано / Есть замечания", ""],
            ["Представитель преподавателей", "", "Согласовано / Есть замечания", ""],
        ],
        [2520, 2160, 2700, 1980],
    )
    for row in doc.tables[-1].rows[1:]:
        set_row_min_height(row, 700)

    doc.add_heading("Разрешение на следующий этап", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(
        "Статус документа:  [  ] Все обязательные решения приняты    "
        "[  ] Остались открытые вопросы    [  ] Нужна новая встреча"
    )
    set_run_font(run, size=10.5, color=INK, bold=True)

    add_answer_box(
        doc,
        "Укажите, разрешено ли переходить к аудиту текущего приложения и проектированию "
        "модели данных. Перечислите любые ограничения или условия."
    )


def set_document_properties(doc):
    props = doc.core_properties
    props.title = "Анкета для согласования финансовой системы Nuriks Academy"
    props.subject = "Решения по тарифам, урокам, платежам, расходам, ролям и отчётности"
    props.author = "Nuriks Academy"
    props.keywords = "финансы, обучение, тарифы, платежи, расходы, зарплата"
    props.comments = "Рабочий документ для обсуждения до разработки"


def main():
    doc = Document()
    configure_styles(doc)
    set_document_properties(doc)

    add_cover(doc)
    add_scope_and_known_values(doc)
    add_questions(doc)
    add_validation_and_signoff(doc)

    # Re-apply section geometry after all content is present.
    configure_sections(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
